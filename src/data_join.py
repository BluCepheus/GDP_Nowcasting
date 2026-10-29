"""Module for unpacking and joining data files."""

import re
import logging
import tempfile
import warnings
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_RAW_DIR = PROJECT_ROOT / "data" / "raw"
LOGGER = logging.getLogger(__name__)
TABLE_PERIOD_PATTERN = re.compile(
    r"(?P<year>\d{4})\s+(?:(?:M(?P<month>\d{1,2}))|"
    r"Q(?P<quarter>[1-4])(?:-Q(?P<quarter_end>[1-4]))?)",
    re.IGNORECASE,
)

MANUFACTURING_LABELS = {
    "Dobra inwestycyjne Investment goods": "Investment goods",
    "Dobra konsumpcyjne Consumer goods": "Consumer goods",
    "Dobra pośrednie Intermediate goods": "Intermediate goods",
    "Energia Energy": "Energy",
    "Nietrwałe dobra konsumpcyjne Non-durable consumer goods": (
        "Non-durable consumer goods"
    ),
    (
        "Produkcja podstawowych substancji farmaceutycznych oraz leków i "
        "pozostałych wyrobów farmaceutycznych Manufacture of basic "
        "pharmaceutical products and pharmaceutical preparations"
    ): "Manufacture of basic pharmaceutical products and pharmaceutical preparations",
    "Trwałe dobra konsumpcyjne Durable consumer goods": "Durable consumer goods",
    "Żywność i napoje Food and beverages": "Food and beverages",
}


ROMAN_TO_MONTH = {
    "I": 1,
    "II": 2,
    "III": 3,
    "IV": 4,
    "V": 5,
    "VI": 6,
    "VII": 7,
    "VIII": 8,
    "IX": 9,
    "X": 10,
    "XI": 11,
    "XII": 12,
}

QUARTERLY_MONTH_GROUPS = (
    (1, 4, 7, 10),
    (2, 5, 8, 11),
    (3, 6, 9, 12),
)


def _parse_consumer_dates(labels: pd.Index) -> pd.DatetimeIndex:
    return pd.to_datetime(
        [
            f"{year}-{ROMAN_TO_MONTH[month]:02d}-01"
            for year, month in (label.split() for label in labels)
        ],
        format="%Y-%m-%d",
    )


def _replace_missing_values(df: pd.DataFrame) -> pd.DataFrame:
    return df.mask(df.eq("x"), pd.NA).mask(df.eq("."), pd.NA)


def _consolidate_ukraine_employment_answers(
    df: pd.DataFrame,
    variable: str,
) -> pd.DataFrame:
    if variable != "The impact of war on employment of workers from Ukraine":
        return df

    old_labels = (
        "minor outflow of workers from Ukraine",
        "serious outflow of workers from Ukraine",
        "minor intflow of workers from Ukraine",
        "serious intflow of workers from Ukraine",
    )
    consolidated_columns = {}
    for column in df.columns:
        for old_label in old_labels:
            suffix = f" - {old_label}"
            if column.endswith(suffix):
                new_label = old_label.split(" ", maxsplit=1)[1]
                new_column = f"{column[:-len(suffix)]} - {new_label}"
                consolidated_columns.setdefault(new_column, []).append(column)
                break

    historical = df.index < pd.Timestamp("2024-01-01")
    for new_column, old_columns in consolidated_columns.items():
        old_values = df[old_columns].sum(axis=1, min_count=1)
        if new_column not in df.columns:
            df[new_column] = pd.NA
        df.loc[historical, new_column] = df.loc[
            historical, new_column
        ].combine_first(old_values.loc[historical])

    return df.drop(
        columns=[
            column
            for columns in consolidated_columns.values()
            for column in columns
        ]
    )


def _align_additional_question_columns(df: pd.DataFrame) -> pd.DataFrame:
    for column in df.columns:
        month_by_quarter = _get_2025_month_by_quarter(df, column)
        if not month_by_quarter:
            continue
        historical = df.loc[df.index.year < 2025, column].dropna()
        for year in sorted(set(historical.index.year)):
            _align_additional_question_year(
                df,
                column,
                year,
                month_by_quarter,
            )
    return df


def _get_2025_month_by_quarter(
    df: pd.DataFrame,
    column: str,
) -> dict[int, int]:
    current_data = df.loc[
        (df.index.year == 2025) & df[column].notna(),
        column,
    ]
    month_by_quarter = {}
    for date in current_data.index:
        previous_month = month_by_quarter.get(date.quarter)
        if previous_month is not None and previous_month != date.month:
            raise ValueError(
                f"Multiple 2025 alignment months for {column!r} "
                f"in quarter {date.quarter}."
            )
        month_by_quarter[date.quarter] = date.month
    return month_by_quarter


def _align_additional_question_year(
    df: pd.DataFrame,
    column: str,
    year: int,
    month_by_quarter: dict[int, int],
) -> None:
    year_data = df.loc[
        (df.index.year == year) & df[column].notna(),
        column,
    ]
    quarter_counts = year_data.groupby(year_data.index.quarter).size()
    is_monthly = (quarter_counts >= 3).any() or (quarter_counts > 1).sum() >= 2
    for quarter in sorted(set(year_data.index.quarter)):
        if quarter not in month_by_quarter:
            continue
        quarter_data = year_data.loc[year_data.index.quarter == quarter]
        aligned_month = month_by_quarter[quarter]
        aligned_data = quarter_data.loc[
            quarter_data.index.month == aligned_month
        ]
        if is_monthly or not aligned_data.empty:
            selected = aligned_data
        else:
            selected = quarter_data.iloc[[-1]]

        df.loc[quarter_data.index, column] = pd.NA
        if not selected.empty:
            target_date = pd.Timestamp(year, aligned_month, 1)
            df.loc[target_date, column] = selected.iloc[-1]


def _split_consumer_data(
    df: pd.DataFrame,
    marker_can_appear_anywhere: bool = False,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    df.columns = [f"Consumer - {column}" for column in df.columns]
    if marker_can_appear_anywhere:
        quarterly_columns = df.columns.str.contains("*", regex=False)
    else:
        quarterly_columns = df.columns.str.endswith("*")
    df_monthly = df.loc[:, ~quarterly_columns].copy()
    df_quarterly = df.loc[:, quarterly_columns].copy()
    df_quarterly.columns = [
        column.replace("*", "")
        if marker_can_appear_anywhere
        else column[:-1]
        for column in df_quarterly.columns
    ]
    df_quarterly.dropna(how="all", inplace=True)
    return df_monthly, df_quarterly


def _load_consumer_indicators(file_path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    df = pd.read_excel(
        file_path,
        sheet_name="CONSUMER TENDENCY INDICATORS",
        skiprows=3,
        header=0,
        index_col=1,
    )[:-3]
    df.drop(columns=["Unnamed: 0"], inplace=True)
    df = df.T
    df.index = _parse_consumer_dates(df.index)
    df = _replace_missing_values(df)
    return _split_consumer_data(df)


def _load_consumer_responses(file_path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    df = pd.read_excel(
        file_path,
        sheet_name="PERCENTAGES OF RESPONSES",
        skiprows=9,
        header=0,
    )[:-5]
    df.iloc[:, 1] = df.iloc[:, 1].ffill().astype("string").str.strip()
    df.iloc[:, 2] = df.iloc[:, 2].astype("string").str.strip()
    df.index = df.iloc[:, 1] + ": " + df.iloc[:, 2]
    df = df.iloc[:, 5:].T
    df.index = _parse_consumer_dates(df.index)
    df = _replace_missing_values(df)
    return _split_consumer_data(df, marker_can_appear_anywhere=True)


def _get_latest_consumer_month(folder_path: Path) -> pd.Timestamp:
    consumer_files = list(folder_path.glob("Consumer*.xlsx"))
    if not consumer_files:
        raise FileNotFoundError(
            f"No Consumer tendency workbook found in {folder_path}."
        )

    observed_dates = []
    for file_path in consumer_files:
        for loader in (_load_consumer_indicators, _load_consumer_responses):
            for frame in loader(file_path):
                observed = frame.dropna(how="all")
                if not observed.empty:
                    observed_dates.append(observed.index.max())
    if not observed_dates:
        raise ValueError("Consumer tendency workbooks have no observations.")
    return max(observed_dates)


def _english_header(value: object) -> str:
    if pd.isna(value):
        return ""
    lines = [line.strip() for line in str(value).splitlines() if line.strip()]
    if not lines:
        return ""
    if len(lines) <= 2:
        english_lines = lines[-1:]
        if len(lines) == 1:
            english_lines = [
                re.sub(r"^.*(?:\t+| {2,})", "", english_lines[0])
            ]
    else:
        english_start = next(
            (
                position
                for position, line in enumerate(lines[1:], start=1)
                if line[0].isupper()
            ),
            len(lines) - 1,
        )
        english_lines = lines[english_start:]
    return re.sub(r"\s+", " ", " ".join(english_lines)).strip()


def _read_main_indicators_sheet(
    file_path: Path,
    sheet_name: str,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Cannot parse header or footer so it will be ignored",
            category=UserWarning,
            module=r"openpyxl\.worksheet\.header_footer",
        )
        workbook = load_workbook(file_path, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        periods = _read_table_periods(worksheet, file_path, sheet_name)
        header_rows = _find_table_header_rows(
            worksheet,
            periods[0][0],
            file_path,
            sheet_name,
        )
        last_column = _find_table_last_column(
            worksheet,
            header_rows,
            periods,
        )
        headers = _read_table_headers(worksheet, header_rows, last_column)
        data = _read_table_data(worksheet, periods, last_column)
    finally:
        workbook.close()

    data = _replace_missing_values(data)
    data = data.mask(data.eq("*"), pd.NA)
    data = data.replace(r"^\s*$", pd.NA, regex=True)
    return headers, data


def _parse_table_period_label(
    value: object,
    file_path: Path,
    sheet_name: str,
) -> pd.Timestamp | None:
    if not isinstance(value, str):
        return None
    match = TABLE_PERIOD_PATTERN.fullmatch(value.strip())
    if match is None:
        return None

    year = int(match.group("year"))
    month = match.group("month")
    if month is not None:
        month_number = int(month)
        if not 1 <= month_number <= 12:
            raise ValueError(
                f"Invalid month in {value!r} in {file_path.name!r}, "
                f"sheet {sheet_name!r}."
            )
        return pd.Timestamp(year, month_number, 1)

    quarter = int(match.group("quarter_end") or match.group("quarter"))
    return pd.Timestamp(year, (quarter - 1) * 3 + 1, 1)


def _read_table_periods(
    worksheet: Worksheet,
    file_path: Path,
    sheet_name: str,
) -> list[tuple[int, pd.Timestamp]]:
    periods = []
    for row_number, (value,) in enumerate(
        worksheet.iter_rows(min_col=1, max_col=1, values_only=True),
        start=1,
    ):
        period_date = _parse_table_period_label(
            value,
            file_path,
            sheet_name,
        )
        if period_date is not None:
            periods.append((row_number, period_date))
    if not periods:
        raise ValueError(
            f"No supported monthly or quarterly observations found in "
            f"{file_path.name!r}, sheet {sheet_name!r}."
        )
    return periods


def _find_table_header_rows(
    worksheet: Worksheet,
    data_start_row: int,
    file_path: Path,
    sheet_name: str,
) -> list[int]:
    header_rows = []
    for row_number in range(1, data_start_row):
        values = (
            worksheet.cell(row_number, column).value
            for column in range(2, worksheet.max_column + 1)
        )
        text_values = [
            value
            for value in values
            if isinstance(value, str) and value.strip()
        ]
        has_header_text = any(
            not value.strip().upper().startswith("TABL")
            for value in text_values
        )
        if has_header_text:
            header_rows.append(row_number)
    if not header_rows:
        raise ValueError(
            f"No column headers found in {file_path.name!r}, "
            f"sheet {sheet_name!r}."
        )
    return header_rows


def _find_table_last_column(
    worksheet: Worksheet,
    header_rows: list[int],
    periods: list[tuple[int, pd.Timestamp]],
) -> int:
    relevant_rows = header_rows + [row for row, _ in periods]
    return max(
        (
            column
            for row_number in relevant_rows
            for column in range(2, worksheet.max_column + 1)
            if worksheet.cell(row_number, column).value is not None
            and str(worksheet.cell(row_number, column).value).strip()
        ),
        default=1,
    )


def _read_table_headers(
    worksheet: Worksheet,
    header_rows: list[int],
    last_column: int,
) -> pd.DataFrame:
    header_values = [
        [
            worksheet.cell(row_number, column).value
            for column in range(2, last_column + 1)
        ]
        for row_number in header_rows
    ]
    header_row_positions = {
        row_number: position
        for position, row_number in enumerate(header_rows)
    }
    for merged_range in worksheet.merged_cells.ranges:
        min_col, min_row, max_col, _ = merged_range.bounds
        if min_row not in header_row_positions or max_col <= min_col:
            continue
        header_row = header_row_positions[min_row]
        first_column = min_col - 2
        end_column = min(max_col - 1, last_column - 1)
        if 0 <= first_column < len(header_values[header_row]):
            header_values[header_row][first_column:end_column] = [
                header_values[header_row][first_column]
            ] * (end_column - first_column)
    return pd.DataFrame(header_values)


def _read_table_data(
    worksheet: Worksheet,
    periods: list[tuple[int, pd.Timestamp]],
    last_column: int,
) -> pd.DataFrame:
    return pd.DataFrame(
        [
            [
                worksheet.cell(row_number, column).value
                for column in range(2, last_column + 1)
            ]
            for row_number, _ in periods
        ],
        index=pd.DatetimeIndex([date for _, date in periods]),
    )


def _quarter_for_consumer_month(
    latest_consumer_month: pd.Timestamp,
) -> tuple[int, int]:
    month_mapping = {
        1: (latest_consumer_month.year - 1, 4),
        4: (latest_consumer_month.year, 1),
        7: (latest_consumer_month.year, 2),
        10: (latest_consumer_month.year, 3),
    }
    try:
        return month_mapping[latest_consumer_month.month]
    except KeyError as error:
        raise ValueError(
            "The latest Consumer tendency month must be January, April, "
            "July, or October."
        ) from error


def _load_quarterly_main_indicators_file(
    file_path: Path,
    latest_consumer_month: pd.Timestamp,
    sheet_name: str = "Kwartalne_Quarterly",
) -> tuple[pd.DataFrame, pd.DataFrame | None]:
    headers, data = _read_main_indicators_sheet(
        file_path,
        sheet_name=sheet_name,
    )
    if file_path.name.startswith("Table 51."):
        data = _decumulate_construction_production(data)

    target_year, target_quarter = _quarter_for_consumer_month(
        latest_consumer_month
    )
    target_date = pd.Timestamp(
        year=target_year,
        month=(target_quarter - 1) * 3 + 1,
        day=1,
    )
    is_target_table = file_path.name.startswith("Table 1.")
    target_column = None
    if is_target_table:
        target_column = _main_indicator_column_name(headers, 0)
        if not target_column:
            raise ValueError(
                f"Missing English header for the target column in "
                f"{file_path.name!r}, sheet {sheet_name!r}."
            )
    target_data = data.loc[data.index == target_date]
    if target_data.empty:
        target = None
        if is_target_table:
            target = data.iloc[:, [0]].copy()
            target.columns = [target_column]
        return pd.DataFrame(index=data.index), target

    target_observations = target_data.iloc[0]
    selected_positions = [
        position
        for position, value in enumerate(target_observations)
        if (not is_target_table or position != 0) and pd.notna(value)
    ]
    columns = [
        _main_indicator_column_name(headers, position)
        for position in selected_positions
    ]
    selected = data.iloc[:, selected_positions].copy()
    selected.columns = columns
    target = None
    if is_target_table:
        target = data.iloc[:, [0]].copy()
        target.columns = [target_column]
    return selected, target


def _decumulate_construction_production(
    data: pd.DataFrame,
) -> pd.DataFrame:
    numeric_data = data.mask(data.eq("(.)"), pd.NA).apply(
        pd.to_numeric,
        errors="raise",
    )
    previous_quarter = numeric_data.copy()
    previous_quarter.index = previous_quarter.index + pd.DateOffset(months=3)
    previous_quarter = previous_quarter.reindex(numeric_data.index)
    decumulated = numeric_data - previous_quarter
    first_quarters = numeric_data.index.quarter == 1
    decumulated.loc[first_quarters] = numeric_data.loc[first_quarters]
    return decumulated


def _main_indicator_column_name(
    headers: pd.DataFrame,
    position: int,
) -> str:
    parts = (
        _english_header(headers.iloc[row, position])
        for row in range(headers.shape[0])
    )
    return " - ".join(part for part in parts if part)


def _group_main_indicator_columns(
    file_path: Path,
    headers: pd.DataFrame,
    data: pd.DataFrame,
    latest_consumer_month: pd.Timestamp,
) -> dict[int, list[tuple[int, str]]]:
    columns_by_lag: dict[int, list[tuple[int, str]]] = {}
    for position in range(data.shape[1]):
        values = data.iloc[:, position]
        last_observation = values.last_valid_index()
        if last_observation is None:
            continue

        column_name = _main_indicator_column_name(headers, position)
        if not column_name:
            raise ValueError(
                f"Missing English header in {file_path.name!r}, "
                f"column {position + 2}."
            )
        month_lag = (
            (latest_consumer_month.year - last_observation.year) * 12
            + latest_consumer_month.month
            - last_observation.month
        )
        if month_lag < 0:
            raise ValueError(
                f"{file_path.name!r}, column {column_name!r} has data from "
                f"{last_observation:%Y-%m}, later than the latest Consumer "
                f"tendency month {latest_consumer_month:%Y-%m}."
            )
        columns_by_lag.setdefault(month_lag, []).append(
            (position, column_name)
        )
    return columns_by_lag


def _load_main_indicators_file(
    file_path: Path,
    latest_consumer_month: pd.Timestamp,
    sheet_name: str = "Miesieczne_Monthly",
) -> dict[int, list[pd.DataFrame]]:
    headers, data = _read_main_indicators_sheet(file_path, sheet_name)
    columns_by_lag = _group_main_indicator_columns(
        file_path,
        headers,
        data,
        latest_consumer_month,
    )
    frames_by_lag = {}
    for month_lag, columns in columns_by_lag.items():
        frame = data.iloc[:, [position for position, _ in columns]].copy()
        frame.columns = [name for _, name in columns]
        frames_by_lag[month_lag] = [frame]
    return frames_by_lag


def _load_construction_file(file_path: Path) -> list[pd.DataFrame]:
    frames = []
    with pd.ExcelFile(file_path) as workbook:
        for sheet_name in workbook.sheet_names:
            if not sheet_name.isdigit():
                continue
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            variable = df.iloc[3, 0]
            df = df.iloc[6:].copy()
            df.columns = [
                column if pd.notna(column) else f"Unnamed: {position}"
                for position, column in enumerate(df.iloc[0])
            ]
            df = df.iloc[1:].set_index(df.columns[0])
            df.index = pd.to_datetime(
                df.index.str.replace("M", "-", regex=False) + "-01",
                format="%Y-%m-%d",
            )
            df.columns = [
                f"Construction - {variable} - {column}"
                for column in df.columns
            ]
            df = _replace_missing_values(df)
            frames.append(df)
    return frames


def _load_additional_questions_file(
    file_path: Path,
) -> list[pd.DataFrame]:
    frames = []
    prefix = file_path.name.split("-", maxsplit=1)[0].title()
    with pd.ExcelFile(file_path) as workbook:
        for sheet_name in workbook.sheet_names:
            if not sheet_name.isdigit():
                continue
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            variable = str(df.iloc[3, 0]).strip()
            last_column = max(
                (
                    position
                    for position in range(1, df.shape[1])
                    if df.iloc[:, position].notna().any()
                ),
                default=0,
            )
            groups = df.iloc[6, 1:last_column + 1].ffill()
            labels = df.iloc[8, 1:last_column + 1]
            columns = []
            for position, (group, label) in enumerate(
                zip(groups, labels),
                start=2,
            ):
                if pd.isna(group) or pd.isna(label):
                    raise ValueError(
                        f"Missing {prefix.lower()} additional-questions header "
                        f"in sheet {sheet_name!r}, column {position}."
                    )
                columns.append(
                    f"{prefix} - {variable} - "
                    f"{_clean_services_header(group)} - "
                    f"{_clean_services_header(label)}"
                )

            data = df.iloc[9:, 1:last_column + 1].copy()
            data.columns = columns
            data.index = pd.to_datetime(
                df.iloc[9:, 0].astype("string").str.strip().str.rstrip(
                    "'"
                ).str.replace("M", "-", regex=False)
                + "-01",
                format="%Y-%m-%d",
            )
            data = _replace_missing_values(data)
            data = _consolidate_ukraine_employment_answers(data, variable)
            data = _align_additional_question_columns(data)
            frames.append(data)
    return frames


def _append_quarterly_frames(
    frames: list[pd.DataFrame],
    quarterly_frames: list[list[pd.DataFrame]],
) -> None:
    for frame in frames:
        alignment_by_column = {}
        for column in frame.columns:
            aligned_months = set(
                frame.index[
                    (frame.index.year == 2025) & frame[column].notna()
                ].month
            )
            if not aligned_months:
                continue
            matching_groups = [
                position
                for position, months in enumerate(QUARTERLY_MONTH_GROUPS)
                if aligned_months.issubset(months)
            ]
            if len(matching_groups) != 1:
                raise ValueError(
                    f"Could not assign {column!r} to a quarterly month group "
                    "based on its 2025 observations."
                )
            alignment_by_column[column] = matching_groups[0]

        for position, (group, months) in enumerate(
            zip(quarterly_frames, QUARTERLY_MONTH_GROUPS)
        ):
            columns = [
                column
                for column, alignment in alignment_by_column.items()
                if alignment == position
            ]
            if columns:
                group.append(
                    frame.loc[
                        frame.index.month.isin(months),
                        columns,
                    ]
                )


def _save_quarterly_files(
    output_dir: Path,
    quarterly_frames: list[list[pd.DataFrame]],
) -> None:
    for position, frames in enumerate(quarterly_frames):
        if not frames:
            continue
        df_quarterly = _concat_unique_columns(
            frames,
            f"quaterly_data_{position}.xlsx",
        )
        df_quarterly.index = (
            df_quarterly.index.to_period("Q-DEC").to_timestamp()
        )
        df_quarterly.to_excel(output_dir / f"quaterly_data_{position}.xlsx")


def _save_monthly_files(
    output_dir: Path,
    monthly_frames: dict[int, list[pd.DataFrame]],
) -> pd.DataFrame:
    df_monthly_0 = pd.DataFrame()
    for position, frames in sorted(monthly_frames.items()):
        if position > 3:
            continue
        df_monthly = (
            _concat_unique_columns(
                frames,
                f"monthly_data_{position}.xlsx",
            )
            if frames
            else pd.DataFrame(index=pd.DatetimeIndex([]))
        )
        if position == 0:
            df_monthly_0 = df_monthly
        df_monthly.to_excel(output_dir / f"monthly_data_{position}.xlsx")
    return df_monthly_0


def _concat_unique_columns(
    frames: list[pd.DataFrame],
    output_name: str,
) -> pd.DataFrame:
    combined = pd.concat(frames, axis=1, join="outer")
    duplicate_columns = combined.columns[
        combined.columns.duplicated(keep=False)
    ].unique()
    if not duplicate_columns.empty:
        counts = pd.Series(combined.columns).value_counts()
        seen: dict[str, int] = {}
        renamed = []
        for column in combined.columns:
            name = str(column)
            seen[name] = seen.get(name, 0) + 1
            if counts[name] == 1:
                renamed.append(name)
            elif seen[name] == 1:
                renamed.append(name)
            else:
                renamed.append(f"{name}__duplicate_{seen[name]}")
        combined = combined.copy()
        combined.columns = renamed
        LOGGER.warning(
            "Duplicate column names in %s were disambiguated by suffixing.",
            output_name,
        )
    return combined


def _publish_output_files(staging_dir: Path, output_dir: Path) -> None:
    staged_files = list(staging_dir.glob("*.xlsx"))
    output_dir.mkdir(parents=True, exist_ok=True)
    for staged_file in staged_files:
        staged_file.replace(output_dir / staged_file.name)

    active_names = {path.name for path in staged_files}
    for pattern in ("monthly_data_*.xlsx", "quaterly_data_*.xlsx"):
        for output_file in output_dir.glob(pattern):
            if output_file.name not in active_names:
                output_file.unlink()
    target_file = output_dir / "target.xlsx"
    if target_file.exists() and "target.xlsx" not in active_names:
        target_file.unlink()


def _load_manufacturing_file(file_path: Path) -> list[pd.DataFrame]:
    frames = []
    with pd.ExcelFile(file_path) as workbook:
        for sheet_name in workbook.sheet_names:
            if not sheet_name.isdigit():
                continue
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            variable = df.iloc[3, 0]
            df = df.iloc[4:, 1:].copy()
            df.columns = [
                column if pd.notna(column) else f"Unnamed: {position}"
                for position, column in enumerate(df.iloc[0])
            ]
            df = df.iloc[1:].copy()
            label_column = df.columns[0]
            df[label_column] = (
                df[label_column]
                .astype("string")
                .str.strip()
                .replace(MANUFACTURING_LABELS)
                .str.replace(r"^.*(?:\t+| {2,})", "", regex=True)
                .str.strip()
            )
            df = df.set_index(label_column).T
            df.index = pd.to_datetime(
                df.index.str.replace("M", "-", regex=False) + "-01",
                format="%Y-%m-%d",
            )
            df.columns = [
                f"Manufacturing - {variable} - {column}"
                for column in df.columns
            ]
            df = _replace_missing_values(df)
            frames.append(df)
    return frames


def _load_services_file(file_path: Path) -> list[pd.DataFrame]:
    frames = []
    with pd.ExcelFile(file_path) as workbook:
        for sheet_name in workbook.sheet_names:
            if not sheet_name.isdigit():
                continue
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            variable = df.iloc[3, 0]
            df = df.iloc[4:, 1:].copy()
            df.columns = [
                column if pd.notna(column) else f"Unnamed: {position}"
                for position, column in enumerate(df.iloc[0])
            ]
            df = df.iloc[1:].copy()
            label_column = df.columns[0]
            df[label_column] = (
                df[label_column]
                .astype("string")
                .str.strip()
                .str.replace(r"[\r\n]+", " ", regex=True)
                .str.replace(r"^.*(?:\t+| {2,})", "", regex=True)
                .str.strip()
            )
            df = df.set_index(label_column).T
            df.index = pd.to_datetime(
                df.index.str.replace("M", "-", regex=False) + "-01",
                format="%Y-%m-%d",
            )
            df.columns = [
                f"Services - {variable} - {column}"
                for column in df.columns
            ]
            df = _replace_missing_values(df)
            frames.append(df)
    return frames


def _clean_services_header(value: object) -> str:
    lines = (
        [line.strip() for line in value.splitlines() if line.strip()]
        if isinstance(value, str)
        else []
    )
    if not lines:
        raise ValueError(f"Empty header value: {value!r}.")
    if len(lines) > 1:
        label = lines[-1]
    else:
        label = re.sub(r"^.*(?:\t+| {2,})", "", lines[0])
    return re.sub(r"\s+", " ", label).strip()


def _parse_services_size_dates(
    date_labels: pd.Series,
    sheet_name: str,
) -> list[pd.Timestamp]:
    dates = list(
        pd.to_datetime(
            date_labels.astype("string").str.replace(
                "M", "-", regex=False
            )
            + "-01",
            format="%Y-%m-%d",
        )
    )
    duplicate_dates = pd.Index(dates).duplicated(keep=False)
    for position in range(1, len(dates) - 1):
        if (
            duplicate_dates[position]
            and dates[position] > dates[position + 1]
        ):
            inferred_date = dates[position - 1] + pd.offsets.MonthBegin(1)
            if (
                dates[position + 1]
                == inferred_date + pd.offsets.MonthBegin(1)
                and inferred_date not in dates
            ):
                dates[position] = inferred_date
    if pd.Index(dates).has_duplicates:
        raise ValueError(
            f"Duplicate dates in services size-class sheet {sheet_name!r}."
        )
    return dates


def _load_services_size_file(file_path: Path) -> list[pd.DataFrame]:
    frames = []
    with pd.ExcelFile(file_path) as workbook:
        for sheet_name in workbook.sheet_names:
            if not sheet_name.isdigit():
                continue
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            variable = df.iloc[3, 0]
            last_column = max(
                (
                    position
                    for position in range(1, df.shape[1])
                    if df.iloc[:, position].notna().any()
                ),
                default=0,
            )
            columns = []
            category = None
            for position in range(1, last_column + 1):
                category_cell = df.iloc[6, position]
                if pd.notna(category_cell) and str(category_cell).strip():
                    category = _clean_services_header(category_cell)
                    columns.append(f"Services - {variable} - {category}")
                    continue

                subcategory_cell = df.iloc[7, position]
                if (
                    category is None
                    or pd.isna(subcategory_cell)
                    or not str(subcategory_cell).strip()
                ):
                    raise ValueError(
                        f"Missing services size-class header in sheet {sheet_name!r}."
                    )
                subcategory = _clean_services_header(subcategory_cell)
                columns.append(
                    f"Services - {variable} - {category} - {subcategory}"
                )

            data = df.iloc[17:, 1:last_column + 1].copy()
            data.columns = columns
            data.index = _parse_services_size_dates(
                df.iloc[17:, 0],
                sheet_name,
            )
            data = _replace_missing_values(data)
            frames.append(data)
    return frames


def _load_trade_file(file_path: Path) -> list[pd.DataFrame]:
    frames = []
    with pd.ExcelFile(file_path) as workbook:
        for sheet_name in workbook.sheet_names:
            if not sheet_name.isdigit():
                continue
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            last_column = max(
                (
                    position
                    for position in range(1, df.shape[1])
                    if df.iloc[:, position].notna().any()
                ),
                default=0,
            )
            sections = df.iloc[4, 1:last_column + 1].ffill()
            columns = []
            for position, (section, label) in enumerate(
                zip(sections, df.iloc[6, 1:last_column + 1]),
                start=1,
            ):
                section_match = re.search(r"\(([^()]*)\)\s*$", str(section))
                if section_match is None or pd.isna(label):
                    raise ValueError(
                        f"Missing trade header in sheet {sheet_name!r}, "
                        f"column {position + 1}."
                    )
                section_name = section_match.group(1).strip()
                label_name = _clean_services_header(label)
                columns.append(f"{section_name} - {label_name}")

            data = df.iloc[7:, 1:last_column + 1].copy()
            data.columns = columns
            data.index = pd.to_datetime(
                df.iloc[7:, 0].astype("string").str.replace(
                    "M", "-", regex=False
                )
                + "-01",
                format="%Y-%m-%d",
            )
            data = _replace_missing_values(data)
            frames.append(data)
    return frames


def _load_business_tendency_file(file_path: Path) -> list[pd.DataFrame]:
    frames = []
    with pd.ExcelFile(file_path) as workbook:
        for sheet_name in workbook.sheet_names:
            if re.fullmatch(r"Tabl\. \d+", sheet_name) is None:
                continue
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            variable = str(df.iloc[1, 0]).replace("\t", "").strip()
            last_column = max(
                (
                    position
                    for position in range(1, df.shape[1])
                    if df.iloc[:, position].notna().any()
                ),
                default=0,
            )
            groups = df.iloc[3, 1:last_column + 1].ffill()
            columns = []
            for position, (group, label) in enumerate(
                zip(groups, df.iloc[4, 1:last_column + 1]),
                start=1,
            ):
                if pd.isna(group) or pd.isna(label):
                    raise ValueError(
                        f"Missing business-tendency header in sheet "
                        f"{sheet_name!r}, column {position + 1}."
                    )
                group_name = _clean_services_header(group)
                label_name = _clean_services_header(label)
                columns.append(
                    f"{variable} - {group_name} - {label_name}"
                )

            data = df.iloc[5:, 1:last_column + 1].copy()
            data.columns = columns
            data.index = pd.to_datetime(
                df.iloc[5:, 0].astype("string").str.replace(
                    "M", "-", regex=False
                )
                + "-01",
                format="%Y-%m-%d",
            )
            data = _replace_missing_values(data)
            frames.append(data)
    return frames


def extract_zips(folder_path: str | Path = DATA_RAW_DIR) -> None:
    """Extract raw and archived ZIP files into the raw target folder."""
    folder_path = Path(folder_path)
    archive_dir = folder_path.parent / "bin"
    found_zips = [
        zip_path
        for source_dir in (folder_path, archive_dir)
        if source_dir.exists()
        for zip_path in source_dir.glob("*.zip")
    ]
    latest_zips: dict[str, Path] = {}
    for zip_path in found_zips:
        group_name = re.sub(
            r"\s+-\s+[A-Za-z]+\s+\d{4}",
            "",
            zip_path.stem,
            flags=re.IGNORECASE,
        )
        current = latest_zips.get(group_name)
        if current is None or zip_path.stat().st_mtime > current.stat().st_mtime:
            latest_zips[group_name] = zip_path

    extracted_from: dict[Path, Path] = {}
    for zip_path in sorted(
        latest_zips.values(),
        key=lambda path: path.stat().st_mtime,
    ):
        with ZipFile(zip_path, "r") as zip_file:
            for member in zip_file.infolist():
                if member.is_dir():
                    continue

                target_path = folder_path / Path(member.filename).name
                previous_source = extracted_from.get(target_path)
                if previous_source is not None:
                    LOGGER.warning(
                        "Archive member collision for %s: %s overwrites %s.",
                        target_path.name,
                        zip_path.name,
                        previous_source.name,
                    )

                with zip_file.open(member) as source:
                    target_path.write_bytes(source.read())
                extracted_from[target_path] = zip_path

        LOGGER.info("Extracted %s to %s.", zip_path.name, folder_path)


def _load_table_workbook(
    file_path: Path,
    latest_consumer_month: pd.Timestamp,
    monthly_frames: dict[int, list[pd.DataFrame]],
    quarterly_frames: list[list[pd.DataFrame]],
) -> pd.DataFrame | None:
    table_match = re.match(r"Table\s+(\d+)\.", file_path.name)
    if table_match is None or int(table_match.group(1)) in {21, 44}:
        return None

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Cannot parse header or footer so it will be ignored",
            category=UserWarning,
            module=r"openpyxl\.worksheet\.header_footer",
        )
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet_names = workbook.sheetnames
    workbook.close()

    target_frame = None
    for sheet_name in sheet_names:
        sheet_name_lower = sheet_name.lower()
        if "monthly" in sheet_name_lower or "miesiecz" in sheet_name_lower:
            for position, frames in _load_main_indicators_file(
                file_path,
                latest_consumer_month,
                sheet_name,
            ).items():
                if position <= 3:
                    monthly_frames.setdefault(position, []).extend(frames)
        elif "quarter" in sheet_name_lower or "kwartal" in sheet_name_lower:
            quarterly_frame, target = _load_quarterly_main_indicators_file(
                file_path,
                latest_consumer_month,
                sheet_name,
            )
            if not quarterly_frame.empty and quarterly_frame.shape[1]:
                quarterly_frames[3].append(quarterly_frame)
            if target is not None:
                target_frame = target
    return target_frame


def _append_monthly_business_file(
    file_path: Path,
    monthly_frames: dict[int, list[pd.DataFrame]],
) -> None:
    loaders = (
        ("bts-construction", _load_construction_file),
        ("bts-manufacturing", _load_manufacturing_file),
        ("bts-services-poland-monthly", _load_services_file),
        ("bts-services-poland-size", _load_services_size_file),
        ("bts-trade", _load_trade_file),
        ("business_tendency_indicators_", _load_business_tendency_file),
    )
    for prefix, loader in loaders:
        if file_path.name.startswith(prefix):
            monthly_frames[0].extend(loader(file_path))
            return


def _append_input_file(
    file_path: Path,
    latest_consumer_month: pd.Timestamp | None,
    monthly_frames: dict[int, list[pd.DataFrame]],
    quarterly_frames: list[list[pd.DataFrame]],
) -> pd.DataFrame | None:
    if file_path.name.startswith("Consumer"):
        for loader in (_load_consumer_indicators, _load_consumer_responses):
            df_monthly, df_quarterly = loader(file_path)
            monthly_frames[0].append(df_monthly)
            quarterly_frames[0].append(df_quarterly)
    elif "additional_questions" in file_path.name:
        _append_quarterly_frames(
            _load_additional_questions_file(file_path),
            quarterly_frames,
        )
    elif file_path.name.startswith("Table "):
        if latest_consumer_month is None:
            raise ValueError(
                "A Consumer tendency workbook is required to route Table "
                f"data from {file_path.name!r}."
            )
        return _load_table_workbook(
            file_path,
            latest_consumer_month,
            monthly_frames,
            quarterly_frames,
        )
    else:
        _append_monthly_business_file(file_path, monthly_frames)
    return None


def join_files(folder_path: str | Path = DATA_RAW_DIR):
    """Join all files into monthly and quarterly data."""

    folder_path = Path(folder_path)
    extract_zips(folder_path)
    file_paths = list(folder_path.glob("*.xlsx"))
    monthly_frames: dict[int, list[pd.DataFrame]] = {0: []}
    quarterly_frames = [[], [], [], []]
    target_frame: pd.DataFrame | None = None
    latest_consumer_month = (
        _get_latest_consumer_month(folder_path)
        if any(path.name.startswith("Table ") for path in file_paths)
        else None
    )

    for file_path in file_paths:
        table_target = _append_input_file(
            file_path,
            latest_consumer_month,
            monthly_frames,
            quarterly_frames,
        )
        if table_target is not None:
            target_frame = table_target

    output_dir = folder_path.parent / "interim"
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=".interim-stage-",
        dir=output_dir.parent,
    ) as staging_path:
        staging_dir = Path(staging_path)
        df_monthly_0 = _save_monthly_files(staging_dir, monthly_frames)
        _save_quarterly_files(staging_dir, quarterly_frames)
        if target_frame is not None:
            target_frame.to_excel(staging_dir / "target.xlsx")
        _publish_output_files(staging_dir, output_dir)

    return df_monthly_0

if __name__ == "__main__":
    join_files()
